"""Общие фикстуры: синтетические документы генерируются прямо в тестах, без реальных данных."""
import html
import io
import zipfile
from collections.abc import Iterator

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app


def make_pdf(pages: list[str]) -> bytes:
    """PDF с текстовым слоем: по одной странице на элемент списка."""
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        body = "".join(f"<p>{html.escape(line)}</p>" for line in text.split("\n"))
        page.insert_htmlbox(pymupdf.Rect(50, 50, 550, 800), body)
    data = doc.tobytes()
    doc.close()
    return data


def _add_scan_page(doc: pymupdf.Document) -> None:
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), False)
    pix.clear_with(200)
    page.insert_image(pymupdf.Rect(50, 50, 450, 250), pixmap=pix)


def make_scanned_pdf(pages: int = 1) -> bytes:
    """PDF без текстового слоя: страницы содержат только картинку (имитация скана)."""
    doc = pymupdf.open()
    for _ in range(pages):
        _add_scan_page(doc)
    data = doc.tobytes()
    doc.close()
    return data


def make_mixed_pdf(text: str, scan_pages: int) -> bytes:
    """Первая страница с текстом, дальше сканы."""
    doc = pymupdf.open(stream=make_pdf([text]), filetype="pdf")
    for _ in range(scan_pages):
        _add_scan_page(doc)
    data = doc.tobytes()
    doc.close()
    return data


def make_encrypted_pdf() -> bytes:
    doc = pymupdf.open(stream=make_pdf(["секретный текст документа"]), filetype="pdf")
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="user", owner_pw="owner")
    doc.close()
    return data


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None, tail: str | None = None) -> bytes:
    """DOCX: абзацы, затем (необязательно) таблица и абзац после неё — для проверки порядка."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    if tail:
        document.add_paragraph(tail)
    return save_docx(document)


def save_docx(document) -> bytes:
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def rewrite_docx_part(data: bytes, part: str, transform) -> bytes:
    """Пересобирает DOCX, применив transform(str) -> str к одной XML-части."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == part:
                content = transform(content.decode("utf-8")).encode("utf-8")
            dst.writestr(item, content)
    return out.getvalue()


def add_docx_entries(data: bytes, entries: dict[str, bytes]) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        for name, content in entries.items():
            dst.writestr(name, content)
    return out.getvalue()


def make_png() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), False)
    pix.clear_with(255)
    return pix.tobytes("png")


def make_zip_without_docx() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", "not a docx")
    return buf.getvalue()


INVOICE_TEXT = "Счёт на оплату № 15 от 12.03.2026\nПоставщик: ООО «Ромашка»\nИНН 7701234567\nИтого: 15 000,00 руб."
INVOICE_TEXT_NORMALIZED = 'Счёт на оплату № 15 от 12.03.2026\nПоставщик: ООО "Ромашка"\nИНН 7701234567\nИтого: 15 000,00 руб.'


@pytest.fixture
def invoice_pdf() -> bytes:
    return make_pdf([INVOICE_TEXT, "Страница 2: условия оплаты"])


@pytest.fixture
def invoice_docx() -> bytes:
    return make_docx(
        ["Договор поставки № 7", "Покупатель: ИП Иванов И. И."],
        table=[["Товар", "Сумма"], ["Бумага А4", "1 200,00"]],
        tail="Подписи сторон",
    )


@pytest.fixture
def scanned_pdf() -> bytes:
    return make_scanned_pdf()


@pytest.fixture
def png_image() -> bytes:
    return make_png()


@pytest.fixture
def template_json() -> str:
    return (
        '{"name": "Счёт", "fields": ['
        '{"name": "inn", "label": "ИНН", "data_type": "inn", "is_required": true, "keywords": ["ИНН"]},'
        '{"name": "comment", "label": "Комментарий", "data_type": "string"}'
        "]}"
    )


@pytest.fixture
def settings() -> Settings:
    return Settings()


@pytest.fixture
def client(settings: Settings) -> Iterator[TestClient]:
    with TestClient(create_app(settings)) as test_client:
        yield test_client


def make_client(**overrides) -> TestClient:
    """Клиент с изменёнными настройками; использовать как контекстный менеджер."""
    return TestClient(create_app(Settings(**overrides)))


def post_extract(client: TestClient, content: bytes, template: str, filename: str = "doc.pdf", headers=None, **form):
    return client.post(
        "/extract",
        files={"file": (filename, content, "application/octet-stream")},
        data={"template": template, **form},
        headers=headers or {},
    )
