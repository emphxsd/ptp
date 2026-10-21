import io
import zipfile

import pytest
from docx import Document

from app.config import Settings
from app.errors import (
    DocumentReadError,
    PasswordProtectedError,
    TextTooLargeError,
    TooManyPagesError,
    UnsafeDocumentError,
)
from app.extraction.file_types import DocumentFormat
from app.extraction.text import extract_text
from tests.conftest import (
    add_docx_entries,
    make_docx,
    make_encrypted_pdf,
    make_mixed_pdf,
    make_pdf,
    make_scanned_pdf,
    rewrite_docx_part,
    save_docx,
)

DEFAULTS = Settings()


# ---------------------------------------------------------------- PDF
def test_pdf_text_is_extracted_from_all_pages(invoice_pdf):
    result = extract_text(invoice_pdf, DocumentFormat.PDF, DEFAULTS)

    assert result.pages == 2
    assert result.has_text_layer is True
    assert result.pages_without_text == 0
    assert "ИНН 7707083893" in result.text
    assert 'ООО "Ромашка"' in result.text  # текст уже нормализован
    assert "условия оплаты" in result.text


def test_scanned_pdf_has_no_text_layer(scanned_pdf):
    result = extract_text(scanned_pdf, DocumentFormat.PDF, DEFAULTS)

    assert result.pages == 1
    assert result.has_text_layer is False
    assert result.text == ""


def test_mixed_pdf_keeps_text_of_text_pages():
    pdf = make_mixed_pdf("Договор аренды № 5 между ООО Альфа и ООО Бета", scan_pages=9)

    result = extract_text(pdf, DocumentFormat.PDF, DEFAULTS)

    assert result.has_text_layer is True
    assert result.pages == 10
    assert result.pages_without_text == 9
    assert "Договор аренды" in result.text


def test_page_text_threshold_is_inclusive():
    exactly_20 = "А" * 20
    nineteen = "Б" * 19

    assert extract_text(make_pdf([exactly_20]), DocumentFormat.PDF, DEFAULTS).has_text_layer is True
    assert extract_text(make_pdf([nineteen]), DocumentFormat.PDF, DEFAULTS).has_text_layer is False


def test_pdf_over_page_limit_is_rejected():
    pdf = make_scanned_pdf(pages=3)

    with pytest.raises(TooManyPagesError):
        extract_text(pdf, DocumentFormat.PDF, Settings(max_pages=2))


def test_pdf_text_over_limit_is_rejected(invoice_pdf):
    with pytest.raises(TextTooLargeError):
        extract_text(invoice_pdf, DocumentFormat.PDF, Settings(max_text_chars=20))


def test_password_protected_pdf_is_reported():
    with pytest.raises(PasswordProtectedError):
        extract_text(make_encrypted_pdf(), DocumentFormat.PDF, DEFAULTS)


def test_broken_pdf_raises_document_read_error():
    with pytest.raises(DocumentReadError):
        extract_text(b"%PDF-1.7 this is not a real pdf", DocumentFormat.PDF, DEFAULTS)


# ---------------------------------------------------------------- DOCX
def test_docx_keeps_order_of_paragraphs_and_tables(invoice_docx):
    result = extract_text(invoice_docx, DocumentFormat.DOCX, DEFAULTS)
    text = result.text

    assert result.has_text_layer is True
    assert result.pages is None
    assert text.index("Договор поставки") < text.index("Бумага А4") < text.index("Подписи сторон")
    assert "Бумага А4 | 1 200,00" in text


def test_docx_keeps_empty_cells_to_preserve_columns():
    docx = make_docx([], table=[["a", "", "c"]])

    assert "a | | c" in extract_text(docx, DocumentFormat.DOCX, DEFAULTS).text


def test_docx_merged_cells_are_not_duplicated():
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "ОБЪЕДИНЕНО"
    table.cell(0, 1).merge(table.cell(0, 2))
    table.cell(0, 1).text = "шапка"
    table.cell(1, 1).text = "x"

    text = extract_text(save_docx(document), DocumentFormat.DOCX, DEFAULTS).text

    assert text.count("ОБЪЕДИНЕНО") == 1
    assert text.count("шапка") == 1


def test_docx_reads_nested_tables_headers_footers_and_content_controls():
    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "ООО Шапка, ИНН 7701234567"
    section.footer.paragraphs[0].text = "Подвал документа"
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).add_table(rows=1, cols=2).rows[0].cells[0].text = "Вложенная ячейка"
    document.add_paragraph("Поле-контрол")
    data = rewrite_docx_part(
        save_docx(document),
        "word/document.xml",
        lambda xml: xml.replace(
            "<w:p><w:r><w:t>Поле-контрол</w:t></w:r></w:p>",
            "<w:sdt><w:sdtContent><w:p><w:r><w:t>Поле-контрол</w:t></w:r></w:p></w:sdtContent></w:sdt>",
        ),
    )
    assert "<w:sdt>" in zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()

    text = extract_text(data, DocumentFormat.DOCX, DEFAULTS).text

    assert "ООО Шапка" in text
    assert "Подвал документа" in text
    assert "Вложенная ячейка" in text
    assert "Поле-контрол" in text


def test_docx_huge_grid_span_does_not_hang():
    docx = make_docx([], table=[["ячейка", "вторая"]])
    patched = rewrite_docx_part(
        docx, "word/document.xml", lambda xml: xml.replace("<w:tcPr>", '<w:tcPr><w:gridSpan w:val="2000000000"/>', 1)
    )

    assert "ячейка" in extract_text(patched, DocumentFormat.DOCX, DEFAULTS).text


def test_docx_with_too_many_table_rows_is_unsafe():
    docx = make_docx([], table=[["r"]] * 6)

    with pytest.raises(UnsafeDocumentError):
        extract_text(docx, DocumentFormat.DOCX, Settings(docx_max_table_rows=5))


def test_docx_with_too_many_archive_entries_is_unsafe():
    docx = add_docx_entries(make_docx(["текст"]), {f"junk/{i}.txt": b"" for i in range(60)})

    with pytest.raises(UnsafeDocumentError):
        extract_text(docx, DocumentFormat.DOCX, Settings(docx_max_entries=50))


def test_docx_zip_bomb_is_unsafe():
    bomb = add_docx_entries(make_docx(["текст"]), {"word/media/bomb.bin": b"\x00" * (8 * 1024 * 1024)})

    with pytest.raises(UnsafeDocumentError):
        extract_text(bomb, DocumentFormat.DOCX, DEFAULTS)


def test_docx_with_broken_xml_raises_document_read_error():
    docx = rewrite_docx_part(make_docx(["текст"]), "word/document.xml", lambda xml: xml[: len(xml) // 2])

    with pytest.raises(DocumentReadError):
        extract_text(docx, DocumentFormat.DOCX, DEFAULTS)


def test_docx_with_wrong_root_element_raises_document_read_error():
    docx = rewrite_docx_part(make_docx(["текст"]), "word/document.xml", lambda _: '<?xml version="1.0"?><root/>')

    with pytest.raises(DocumentReadError):
        extract_text(docx, DocumentFormat.DOCX, DEFAULTS)


def test_broken_docx_raises_document_read_error():
    with pytest.raises(DocumentReadError):
        extract_text(b"PK\x03\x04 broken zip", DocumentFormat.DOCX, DEFAULTS)


def test_empty_docx_has_no_text():
    result = extract_text(make_docx([]), DocumentFormat.DOCX, DEFAULTS)

    assert result.text == ""
    assert result.has_text_layer is False


# ---------------------------------------------------------------- изображения
def test_image_has_no_text_layer(png_image):
    result = extract_text(png_image, DocumentFormat.PNG, DEFAULTS)

    assert result.has_text_layer is False
    assert result.pages == 1
