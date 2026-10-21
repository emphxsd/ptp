"""Вывод синтетического документа в PDF с текстовым слоем, DOCX и «скан» (изображение без текста).
Скан имитирует сканер: пониженное разрешение, небольшой наклон, шум и сжатие JPEG."""
import html
import io
import random

import numpy as np
import pymupdf
from docx import Document

from evaluation.documents import Block, SyntheticDocument

PAGE = pymupdf.paper_rect("a4")
MARGIN = 56
CSS = (
    "body {font-family: sans-serif; font-size: 11pt;} h1, h2, h3 {font-size: 13pt;} "
    "table {border-collapse: collapse;} td {border: 1px solid #000; padding: 2px 4px;}"
)
SCAN_DPI = (150, 200)
SCAN_MAX_ANGLE = 1.2  # градусов
SCAN_NOISE_SIGMA = 10.0
SCAN_JPEG_QUALITY = (60, 85)


def to_html(blocks: tuple[Block, ...]) -> str:
    parts = []
    for kind, content in blocks:
        if kind == "h":
            parts.append(f"<h3>{html.escape(str(content))}</h3>")
        elif kind == "p":
            parts.append(f"<p>{html.escape(str(content))}</p>")
        else:
            rows = "".join(
                "<tr>" + "".join(f"<td>{html.escape(cell)}</td>" for cell in row) + "</tr>" for row in content
            )
            parts.append(f"<table>{rows}</table>")
    return "".join(parts)


def to_pdf(document: SyntheticDocument) -> bytes:
    """Многостраничная вёрстка через Story: длинный договор переносится на следующие страницы, а не сжимается."""
    story = pymupdf.Story(html=to_html(document.blocks), user_css=CSS)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    where = PAGE + (MARGIN, MARGIN, -MARGIN, -MARGIN)
    more = True
    while more:
        device = writer.begin_page(PAGE)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buffer.getvalue()


def to_docx(document: SyntheticDocument) -> bytes:
    doc = Document()
    for kind, content in document.blocks:
        if kind == "h":
            doc.add_heading(str(content), level=2)
        elif kind == "p":
            doc.add_paragraph(str(content))
        else:
            rows = list(content)
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def to_scan(document: SyntheticDocument, rng: random.Random, as_image: bool = False) -> bytes:
    """PDF из картинок страниц (или JPEG первой страницы) — текстового слоя нет, нужен OCR."""
    source = pymupdf.open(stream=to_pdf(document), filetype="pdf")
    images = [_scan_page(page, rng) for page in source]
    source.close()
    if as_image:
        return images[0]
    out = pymupdf.open()
    for image in images:
        page = out.new_page(width=PAGE.width, height=PAGE.height)
        page.insert_image(page.rect, stream=image)
    data = out.tobytes(deflate=True)
    out.close()
    return data


def _scan_page(page: pymupdf.Page, rng: random.Random) -> bytes:
    zoom = rng.randint(*SCAN_DPI) / 72
    matrix = pymupdf.Matrix(zoom, zoom).prerotate(rng.uniform(-SCAN_MAX_ANGLE, SCAN_MAX_ANGLE))
    pixmap = page.get_pixmap(matrix=matrix, colorspace=pymupdf.csGRAY, alpha=False)
    pixels = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(pixmap.height, pixmap.stride)[:, : pixmap.width]
    noise = np.random.default_rng(rng.getrandbits(32)).normal(0, SCAN_NOISE_SIGMA, pixels.shape)
    noisy = np.clip(pixels.astype(np.float32) * rng.uniform(0.9, 1.0) + noise, 0, 255).astype(np.uint8)
    result = pymupdf.Pixmap(pymupdf.csGRAY, pixmap.width, pixmap.height, noisy.tobytes(), False)
    return result.tobytes("jpeg", jpg_quality=rng.randint(*SCAN_JPEG_QUALITY))
